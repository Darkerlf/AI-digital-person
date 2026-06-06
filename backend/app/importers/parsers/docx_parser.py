from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


def read_docx_tables(path: Path) -> list[list[str]]:
    document = Document(str(path))
    rows: list[list[str]] = []
    for table in document.tables:
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
    return rows


def read_docx_paragraphs(path: Path) -> str:
    document = Document(str(path))
    return "\n".join(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())


def iter_document_blocks(document: DocxDocument):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def read_docx_blocks(path: Path) -> str:
    document = Document(str(path))
    blocks: list[str] = []
    for block in iter_document_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                blocks.append(text)
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
    return "\n".join(blocks)
