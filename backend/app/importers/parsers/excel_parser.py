from pathlib import Path

from openpyxl import load_workbook


def read_excel_rows(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return []
    header = [str(item).strip() if item is not None else "" for item in header_row]
    return [dict(zip(header, row, strict=False)) for row in rows]
