from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.importers.behavior_excel_importer import BehaviorExcelImporter
from app.importers.knowledge_doc_importer import KnowledgeDocImporter
from app.importers.scenic_docx_importer import ScenicDocxImporter
from app.main import app
from app.models.knowledge_document import KnowledgeDocument
from app.models.scenic_area import ScenicArea


def test_scenic_docx_importer_extracts_table_rows() -> None:
    importer = ScenicDocxImporter()
    rows = importer.parse(Path("tests/fixtures/scenic_spots.docx"))

    assert rows[0]["spot_code"] == "LS-001"
    assert rows[0]["name"] == "灵山大照壁"


def test_knowledge_doc_importer_includes_docx_table_content(tmp_path: Path) -> None:
    path = tmp_path / "guide.docx"
    document = Document()
    document.add_paragraph("灵山胜境导览")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "说明"
    table.rows[1].cells[0].text = "门票"
    table.rows[1].cells[1].text = "成人票参考价210元"
    document.save(path)

    payload = KnowledgeDocImporter().parse(path)

    assert "灵山胜境导览" in payload["content_text"]
    assert "门票" in payload["content_text"]
    assert "成人票参考价210元" in payload["content_text"]


def test_scenic_docx_importer_skips_header_rows_in_multiple_tables(tmp_path: Path) -> None:
    path = tmp_path / "spots.docx"
    headers = [
        "景区名称",
        "景点ID",
        "景点名称",
        "具体位置",
        "建筑/景观参数",
        "核心功能",
        "文化内涵",
        "详细介绍",
        "游玩亮点",
        "演艺/开放信息",
        "备注",
    ]
    document = Document()
    for area, code, name in [("灵山胜境", "LS-001", "灵山大照壁"), ("拈花湾禅意小镇", "NH-001", "拈花广场")]:
        table = document.add_table(rows=2, cols=len(headers))
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header
        values = [area, code, name, "位置", "参数", "功能", "文化", "介绍", "亮点", "开放信息", "备注"]
        for index, value in enumerate(values):
            table.rows[1].cells[index].text = value
    document.save(path)

    rows = ScenicDocxImporter().parse(path)

    assert [row["spot_code"] for row in rows] == ["LS-001", "NH-001"]
    assert all(row["spot_code"] != "景点ID" for row in rows)


def test_behavior_excel_importer_normalizes_tourism_dataset_schema(tmp_path: Path) -> None:
    path = tmp_path / "behavior.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "tourist_id",
            "user_nickname",
            "attraction_name",
            "attraction_content",
            "attraction_type",
            "visit_date",
            "stay_duration",
            "ticket_cost",
            "total_cost",
            "satisfaction",
        ]
    )
    sheet.append(["u1", "游客A", "灵山大佛", "文化景点", "佛教文化", "2026-04-07", 90, 210, 310, 4.8])
    workbook.save(path)

    rows = BehaviorExcelImporter().parse(path)
    summary = BehaviorExcelImporter().build_knowledge_summary(path)

    assert rows[0]["event_time"] == "2026-04-07"
    assert rows[0]["event_type"] == "visit"
    assert rows[0]["spot_name"] == "灵山大佛"
    assert rows[0]["event_value"] == "90"
    assert rows[0]["ticket_cost"] == 210
    assert "灵山大佛" in summary
    assert "平均门票消费 210.00 元" in summary
